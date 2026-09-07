"""Excel reading layer (port of Swift Services/ExcelParser.swift).

openpyxl differs from CoreXLSX in cell semantics:
  - shared strings and rich text are already collapsed to plain ``str``;
  - date-formatted numeric cells come back as ``datetime`` objects;
  - numeric cells come back as ``int``/``float`` instead of the raw XML text;
  - merged regions expose their value only at the top-left cell (other cells
    are ``MergedCell`` and must not be read directly).

To match Swift's *final string* output we normalize on read:
  - datetime  -> "yyyy/M/d" (no leading zeros), matching Swift's DateFormatter;
  - int       -> decimal string without a trailing ".0";
  - float     -> shortest repr (openpyxl only keeps non-integral floats);
  - merged    -> top-left value copied into empty cells of the region.
"""

from __future__ import annotations

import datetime

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.worksheet.worksheet import Worksheet

from . import config, console, errors
from .excel_columns import letter as col_letter
from .models import ExcelContext


def _cell_to_string(value: object) -> str:
    """Reproduce the string a Swift CoreXLSX read would yield for a cell."""
    if value is None:
        return ""
    if isinstance(value, datetime.datetime):
        # Swift DateFormatter "yyyy/M/d" — no zero padding, date part only.
        return f"{value.year}/{value.month}/{value.day}"
    if isinstance(value, datetime.date):
        return f"{value.year}/{value.month}/{value.day}"
    if isinstance(value, datetime.time):
        # openpyxl parses time-only formatted serials into datetime.time.
        # Swift treats them as OLE dates: date part of (1899-12-30 + serial).
        # Recover the serial from the time-of-day fraction, like CoreXLSX does.
        serial = (
            value.hour * 3600 + value.minute * 60 + value.second
        ) / 86400.0
        base = datetime.date(1899, 12, 30)
        d = base + datetime.timedelta(days=serial)
        return f"{d.year}/{d.month}/{d.day}"
    if isinstance(value, bool):
        return "1" if value else "0"
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        return _float_to_string(value)
    # str and any other scalar: use as-is.
    return str(value)


def _float_to_string(value: float) -> str:
    # openpyxl yields int for integral floats; a real float needs shortest repr
    # that reconstructs it (matching the decimal Excel/XML stored).
    if value.is_integer():
        return str(int(value))
    return repr(value)


def get_headers_dict(grid: list[list[str]], max_col: int) -> dict[str, list[int]]:
    """Build the header lookup (port of Swift getHeadersDict).

    Keys are the column letter, its lowercase, and the header text; each maps
    to the list of 0-based column indices carrying that key.
    """
    result: dict[str, list[int]] = {}
    if not grid:
        return result
    first_row = grid[0]
    for i in range(max_col):
        if i >= len(first_row):
            break
        header = first_row[i]
        letter_ = col_letter(i)
        if letter_ is None:
            continue
        result.setdefault(letter_, []).append(i)
        result.setdefault(letter_.lower(), []).append(i)
        result.setdefault(header, []).append(i)
    return result


def _abort(msg: str) -> errors.AbortError:
    return errors.AbortError(msg)


def open_workbook(path: str):
    """Load a workbook (data_only=True), raising ExcelParsingError on failure."""
    try:
        return load_workbook(path, data_only=True, read_only=False)
    except Exception as exc:  # openpyxl raises various archive/parse errors
        raise errors.ExcelParsingError(exc) from exc


def sheet_names(workbook) -> list[str]:
    """Return worksheet names in natural (file) order."""
    return list(workbook.sheetnames)


def collect_cells(ws: Worksheet) -> tuple[dict[tuple[int, int], str], int]:
    """Collect every cell's resolved string plus the merged-region fill.

    Returns (data, max_row) where data is keyed by (row_1based, col_0based)
    and reflects Swift's read+merge-fill semantics (see module docstring).
    Shared with the IPO processor which reads raw rows directly.
    """
    data: dict[tuple[int, int], str] = {}

    # openpyxl exposes merged-region value only at the top-left Cell; every
    # other cell in the region is a MergedCell, skipped below. The top-left
    # value is therefore collected naturally and the whole region is filled
    # afterwards from it.
    merged_ranges = list(ws.merged_cells.ranges)

    for row in ws.iter_rows():
        for cell in row:
            if isinstance(cell, MergedCell):
                continue
            v = cell.value
            if v is None:
                continue
            data[(cell.row, cell.column - 1)] = _cell_to_string(v)

    # --- Apply merged-region fill from each region's top-left value. ------
    for mr in merged_ranges:
        tl_row = mr.min_row
        tl_col0 = mr.min_col - 1
        n_rows = mr.max_row - mr.min_row + 1
        n_cols = mr.max_col - mr.min_col + 1
        area = n_rows * n_cols
        if area > config.Limits.max_merge_cell_area:
            console.warning(
                f"Skipping abnormally large merged cell area ({area} cells). "
                "Possible corrupted Excel format."
            )
            continue

        top_left = data.get((tl_row, tl_col0), "")
        for r in range(mr.min_row, mr.max_row + 1):
            for c0 in range(mr.min_col - 1, mr.max_col):
                key = (r, c0)
                cur = data.get(key)
                if cur is None or cur == "":
                    data[key] = top_left

    max_row = max((r for (r, _) in data), default=0)
    return data, max_row


def read_excel_to_grid(
    path: str,
    sheet_name: str,
    start_row: int = config.Limits.default_excel_start_row,
    workbook=None,
) -> ExcelContext:
    """Read a worksheet into an ExcelContext (port of readExcelToGrid).

    If `workbook` is provided it is reused (caller owns closing it); otherwise
    a workbook is opened for this call.
    """
    if start_row < config.Limits.min_data_start_row:
        raise _abort(
            f"Start row must be at least {config.Limits.min_data_start_row} "
            f"(Row {config.Limits.header_row} is reserved for headers)."
        )

    wb = workbook
    close_after = workbook is None
    if wb is None:
        wb = open_workbook(path)

    try:
        matched = next(
            (n for n in wb.sheetnames if n.casefold() == sheet_name.casefold()),
            None,
        )
        if matched is None:
            raise _abort(f"Sheet '{sheet_name}' not found in the Excel file.")
        ws: Worksheet = wb[matched]
        return _read_worksheet(ws, path, matched, start_row)
    finally:
        if close_after:
            wb.close()


def _read_worksheet(
    ws: Worksheet, path: str, sheet_name: str, start_row: int
) -> ExcelContext:
    data, max_row = collect_cells(ws)

    # --- Row count safety check. ------------------------------------------
    if max_row > config.Limits.max_excel_rows:
        raise _abort(
            f"Excel row count ({max_row}) exceeds safe limit "
            f"({config.Limits.max_excel_rows})."
        )
    if max_row < start_row:
        raise _abort(
            f"Sheet '{sheet_name}' has no data rows (max row {max_row} is less "
            f"than start row {start_row})."
        )

    # --- Determine the data width (columns). ------------------------------
    max_col = 0
    for r in range(start_row, max_row + 1):
        row_cols = [c for (rr, c) in data if rr == r]
        if row_cols:
            max_col = max(max_col, max(row_cols) + 1)
    if max_col == 0:
        header_cols = [c for (rr, c) in data if rr == start_row - 1]
        if header_cols:
            max_col = max(header_cols) + 1
    max_col = min(max_col, config.Limits.max_excel_columns)

    # --- Build the grid. ---------------------------------------------------
    grid: list[list[str]] = []
    # Header row (Excel row startRow - 1).
    h_row: list[str] = []
    for c in range(max_col):
        letter_ = col_letter(c)
        val = data.get((start_row - 1, c))
        if letter_ is None or val is None:
            h_row.append(f"{config.Defaults.default_column_prefix}{c + 1}")
        else:
            h_row.append(val)
    grid.append(h_row)

    # Data rows.
    for r in range(start_row, max_row + 1):
        row: list[str] = []
        for c in range(max_col):
            row.append(data.get((r, c), ""))
        grid.append(row)

    return ExcelContext(
        file_path=path,
        sheet_name=sheet_name,
        start_row=start_row,
        headers=get_headers_dict(grid, max_col),
        grid=grid,
    )
