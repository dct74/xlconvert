"""Phase 1 tests: Excel reading layer (excel_parser.py)."""

import datetime

import pytest
from openpyxl import Workbook

from exconverter import excel_parser
from exconverter.errors import AbortError, ExcelParsingError


@pytest.fixture()
def basic_workbook(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Probe"
    ws["A1"] = "名称"
    ws["B1"] = "日期"
    ws["C1"] = "数量"
    ws["D1"] = "备注"
    ws["A2"] = "项目A"
    ws["B2"] = datetime.datetime(2026, 1, 5)
    ws["B2"].number_format = "yyyy/m/d"
    ws["A3"] = "项目B"
    ws["B3"] = 46037
    ws["B3"].number_format = "yyyy/m/d"  # serial -> 2026/1/15
    ws["C2"] = 46037  # int
    ws["C3"] = 3.14  # float
    ws["D2"] = "合并备注"
    ws.merge_cells("D2:D3")
    ws["E1"] = "备注头合并"
    ws.merge_cells("E1:F1")
    ws["F2"] = 123
    path = tmp_path / "basic.xlsx"
    wb.save(path)
    return str(path)


def test_grid_shape_and_values(basic_workbook):
    ctx = excel_parser.read_excel_to_grid(basic_workbook, "Probe", 2)
    assert ctx.sheet_name == "Probe"
    assert ctx.start_row == 2
    assert ctx.grid[0] == ["名称", "日期", "数量", "备注", "备注头合并", "备注头合并"]
    # Data row at Excel row 2 -> grid index 1.
    assert ctx.grid[1] == ["项目A", "2026/1/5", "46037", "合并备注", "", "123"]
    # Data row at Excel row 3 -> grid index 2.
    assert ctx.grid[2] == ["项目B", "2026/1/15", "3.14", "合并备注", "", ""]


def test_header_dict_keys(basic_workbook):
    ctx = excel_parser.read_excel_to_grid(basic_workbook, "Probe", 2)
    h = ctx.headers
    assert h["A"] == [0]
    assert h["a"] == [0]
    assert h["名称"] == [0]
    assert h["B"] == [1]
    assert h["日期"] == [1]
    assert h["C"] == [2]


def test_date_no_leading_zero(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "D"
    ws["A1"] = "日期"
    # day/month 5/1 must not become 05/01
    ws["A2"] = datetime.datetime(2026, 5, 1)
    ws["A2"].number_format = "yyyy/m/d"
    ws["B1"] = "中文"
    ws["B2"] = datetime.datetime(2026, 2, 3)
    ws["B2"].number_format = 'yyyy"年"m"月"d"日"'
    path = tmp_path / "d.xlsx"
    wb.save(path)
    ctx = excel_parser.read_excel_to_grid(str(path), "D", 2)
    assert ctx.grid[1] == ["2026/5/1", "2026/2/3"]


def test_numeric_and_text_stay_raw(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "N"
    ws["A2"] = "123"  # text
    ws["B2"] = 123  # int
    ws["C2"] = 123.0  # integral float -> int
    ws["D2"] = 3.5  # float
    path = tmp_path / "n.xlsx"
    wb.save(path)
    ctx = excel_parser.read_excel_to_grid(str(path), "N", 2)
    assert ctx.grid[1] == ["123", "123", "123", "3.5"]


def test_empty_header_default_name(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "H"
    ws["A1"] = "序号"
    ws["B2"] = "x"  # no header at B1
    path = tmp_path / "h.xlsx"
    wb.save(path)
    ctx = excel_parser.read_excel_to_grid(str(path), "H", 2)
    assert ctx.grid[0] == ["序号", "Column_2"]


def test_case_insensitive_sheet_name(basic_workbook):
    ctx = excel_parser.read_excel_to_grid(basic_workbook, "probe", 2)
    assert ctx.sheet_name == "Probe"


def test_sheet_not_found(basic_workbook):
    with pytest.raises(AbortError, match="not found"):
        excel_parser.read_excel_to_grid(basic_workbook, "Nope", 2)


def test_start_row_too_small(basic_workbook):
    with pytest.raises(AbortError, match="reserved for headers"):
        excel_parser.read_excel_to_grid(basic_workbook, "Probe", 1)


def test_not_an_xlsx(tmp_path):
    bad = tmp_path / "bad.xlsx"
    bad.write_bytes(b"this is not a zip archive")
    with pytest.raises(ExcelParsingError):
        excel_parser.read_excel_to_grid(str(bad), "Probe", 2)


def test_time_only_cell_does_not_crash(tmp_path):
    """Regression: openpyxl parses time-only (hh:mm:ss) cells into
    datetime.time; reading must not raise AttributeError (no .year)."""
    wb = Workbook()
    ws = wb.active
    ws.title = "T"
    ws["A1"] = "时间"
    ws["A2"] = datetime.time(9, 15, 0)
    ws["A2"].number_format = "hh:mm:ss"
    path = tmp_path / "t.xlsx"
    wb.save(path)
    ctx = excel_parser.read_excel_to_grid(str(path), "T", 2)
    # time-only serial is sub-1-day, so date part stays at the CoreXLSX epoch.
    assert ctx.grid[1] == ["1899/12/30"]


def test_time_serial_matches_corexlsx_epoch():
    """Time serial fraction -> 1899-12-30 + fraction day -> 1899/12/30."""
    assert excel_parser._cell_to_string(datetime.time(9, 15, 0)) == "1899/12/30"
    assert excel_parser._cell_to_string(datetime.time(0, 0, 0)) == "1899/12/30"
